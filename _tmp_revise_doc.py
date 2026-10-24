"""Temporary: add try-on and refresh features to the paper (deleted after use)."""
import copy
import shutil

from docx import Document
from docx.text.paragraph import Paragraph

PATH = r"C:\Folder Kuliah\Semester 5\KECERDASAN BUATAN\Makalah kecerdasan buatan.docx"
BACKUP = r"C:\Folder Kuliah\Semester 5\KECERDASAN BUATAN\Makalah kecerdasan buatan (backup sebelum revisi 2).docx"

open(PATH, "r+b").close()  # fail fast if Word still has the file open
shutil.copyfile(PATH, BACKUP)

d = Document(PATH)
P = list(d.paragraphs)

# Guard: make sure we are editing the expected layout.
assert P[80].text.startswith("3.4 Alur Kerja"), P[80].text
assert P[92].text.startswith("3.7 Bentuk Keluaran"), P[92].text
assert P[127].text.startswith("5.1 Kesimpulan"), P[127].text
assert "(gemini-flash-latest)." in P[53].text, P[53].text
assert P[85].text.startswith("Aplikasi memvalidasi keluaran"), P[85].text
assert P[87].text.startswith("Jika Gemini gagal"), P[87].text


def set_text(par, text):
    runs = par.runs
    if not runs:
        par.add_run(text)
        return
    runs[0].text = text
    for r in runs[1:]:
        r._element.getparent().remove(r._element)


def insert_after(par, text, template=None):
    src = template if template is not None else par
    el = copy.deepcopy(src._p)
    par._p.addnext(el)
    new = Paragraph(el, par._parent)
    set_text(new, text)
    return new


def add_row(table, cells, clone_index=-1):
    tr = copy.deepcopy(table.rows[clone_index]._tr)
    table._tbl.append(tr)
    row = table.rows[-1]
    for cell, text in zip(row.cells, cells):
        cell.text = text


# ---------------- 2.3 Gemini: backup models
set_text(P[53], P[53].text.replace(
    "(gemini-flash-latest).",
    "(gemini-flash-latest), dengan model cadangan gemini-3.1-flash-lite dan "
    "gemini-3.5-flash-lite yang dicoba berurutan apabila model utama melewati batas waktu "
    "atau sedang sibuk.",
))

# ---------------- 3.4 steps: accessories only on request, smarter fallback
set_text(P[85], P[85].text.rstrip() + (
    " Topi dan kacamata hanya disertakan apabila prompt pengguna memintanya."
))
set_text(P[87], (
    "Jika seluruh model Gemini gagal atau tidak menghasilkan pilihan valid, fallback berbasis "
    "aturan memilih atasan, bawahan, dan alas kaki yang paling cocok dengan kata kunci prompt, "
    "dengan gender yang konsisten, bukan produk anak kecuali diminta, dan tetap masuk anggaran."
))

# ---------------- 3.4 Alur kerja: slot completion + refresh
s1 = insert_after(P[86], (
    "Jika slot inti (atasan, bawahan, atau alas kaki) belum terisi, aplikasi melengkapinya "
    "dengan produk dari kandidat atau katalog yang masih masuk sisa anggaran."
))
s2 = insert_after(P[87], (
    "Jika pengguna meminta opsi lain (refresh), kombinasi yang sudah pernah ditampilkan "
    "dikirim kembali; aplikasi memastikan hasil baru tidak sama persis dengan kombinasi "
    "tersebut (lihat subbab 3.9)."
))
_ = (s1, s2)

# ---------------- 3.6 Endpoint table
t0 = d.tables[0]
t0.rows[-1].cells[1].text = (
    "Menghasilkan rekomendasi outfit dan pratinjau (parameter: budget, prompt, mode, image, "
    "refresh, previous_combinations)."
)
add_row(t0, ["POST /api/v1/outfits/try-on",
             "Membuat pratinjau dari item pilihan pengguna per slot, dengan validasi kategori."])

# ---------------- 3.7 Output: combination + notice
set_text(P[93], P[93].text.rstrip() + (
    " Respons juga memuat field combination, yaitu daftar ID item pada outfit tersebut, yang "
    "digunakan saat meminta opsi lain, serta field notice apabila tidak ada kombinasi baru "
    "yang dapat dibentuk."
))

# ---------------- New 3.8 and 3.9
h38 = insert_after(P[93], "3.8 Uji Outfit Manual (Try-On)", template=P[92])
p38 = insert_after(h38, (
    "Selain rekomendasi otomatis, pengguna dapat menyusun outfit sendiri melalui endpoint "
    "POST /api/v1/outfits/try-on dengan mengisi ID item pada lima slot: atasan_id, bawahan_id, "
    "alas_kaki_id, kacamata_id, dan topi_id. Minimal satu slot wajib diisi. Setiap ID "
    "divalidasi terhadap katalog dan kategorinya harus sesuai slot: atasan menerima kategori "
    "ber-body_zone upper atau outer, bawahan menerima lower, alas kaki menerima feet sehingga "
    "kategori sepatu dan sendal sama-sama diterima, sedangkan kacamata dan topi hanya menerima "
    "kategori dengan nama yang sama. Jika ID tidak ditemukan atau kategorinya tidak sesuai "
    "(misalnya item atasan diisikan pada kacamata_id), aplikasi menolak permintaan dengan "
    "status 422 dan menampilkan seluruh kesalahan sekaligus. Jika valid, aplikasi menghitung "
    "total harga dari basis data dan membuat pratinjau komposit tanpa memanggil Gemini."
), template=P[93])
h39 = insert_after(p38, "3.9 Refresh Rekomendasi", template=P[92])
insert_after(h39, (
    "Apabila pengguna tidak menyukai opsi yang ditampilkan, frontend mengirim ulang "
    "permintaan rekomendasi dengan refresh=true dan previous_combinations berisi kombinasi "
    "yang sudah pernah ditampilkan. Pada refresh, kandidat shortlist diambil secara acak, "
    "temperatur model dinaikkan, dan Gemini diminta tidak mengulang kombinasi tersebut. "
    "Karena keluaran model tidak dapat dijamin, aplikasi juga memeriksa hasil akhirnya: jika "
    "himpunan ID item sama dengan salah satu kombinasi sebelumnya, aplikasi mengganti item "
    "satu slot demi satu slot (dimulai dari atasan) dengan produk lain pada slot yang sama "
    "yang tetap masuk anggaran. Sebagian item boleh sama dengan opsi sebelumnya, misalnya "
    "bawahan sama tetapi atasan berbeda, asalkan kombinasinya berbeda. Jika tidak ada "
    "kombinasi baru yang masuk anggaran, aplikasi tetap mengembalikan hasil dan mengisi "
    "field notice."
), template=P[93])

# ---------------- 4.1 Test results table
t1 = d.tables[1]
add_row(t1, ["Uji outfit manual",
             "Pratinjau dibuat dari item pilihan; sepatu dan sendal diterima sebagai alas kaki.",
             "Berhasil: lima slot terisi menghasilkan pratinjau (200); sepatu dan sendal "
             "diterima pada alas_kaki_id."])
add_row(t1, ["Validasi slot kategori",
             "Item dengan kategori yang tidak sesuai slot ditolak.",
             "Berhasil: item atasan pada kacamata_id ditolak (422) dengan pesan kesalahan; "
             "beberapa kesalahan dilaporkan sekaligus."])
add_row(t1, ["Refresh rekomendasi",
             "Kombinasi baru tidak sama dengan kombinasi sebelumnya.",
             "Berhasil: empat kali refresh berturut-turut menghasilkan kombinasi baru, termasuk "
             "saat respons AI (disimulasikan) selalu mengulang kombinasi yang sama; anggaran "
             "sangat kecil memunculkan notice."])

# ---------------- BAB V
set_text(P[129], (
    "Sistem menekankan bahwa keluaran AI tidak langsung dianggap benar. Produk divalidasi "
    "terhadap katalog, sedangkan total harga dihitung dari basis data. Pengujian fungsional "
    "menunjukkan fitur kategori, katalog, rekomendasi, validasi ID, kepatuhan anggaran, "
    "fallback, pratinjau komposit, uji outfit manual dengan validasi kategori per slot, dan "
    "refresh rekomendasi berjalan sesuai harapan. Meski demikian, relevansi rekomendasi dan "
    "usability belum dapat disimpulkan sebelum evaluasi oleh responden dilakukan."
))

d.save(PATH)
print("saved; backup:", BACKUP)
