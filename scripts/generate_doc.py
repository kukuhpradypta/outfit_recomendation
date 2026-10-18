"""Generate the concept/journal document aligned to the current backend.

Produces a .docx describing the outfit recommendation feature as actually
implemented: Gemini multimodal (Vision + reasoning) selects products from the
real catalog with an anti-hallucination guarantee and a deterministic budget
fallback; outfit preview is a free local Pillow composite (mannequin / photo).

Run:
    python scripts/generate_doc.py
"""
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt, RGBColor

OUTPUT = Path(
    r"C:\Folder Kuliah\Semester 5\KECERDASAN BUATAN"
    r"\Fitur Rekomendasi Pakaian Berbasis Budget dan Acara.docx"
)


def add_bullets(doc: Document, items: list[str]) -> None:
    for text in items:
        doc.add_paragraph(text, style="List Bullet")


def build() -> Document:
    doc = Document()

    # Base font
    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(11)

    # Title
    doc.add_heading("Fitur Rekomendasi Pakaian Berbasis Budget dan Acara", level=0)
    sub = doc.add_paragraph(
        "Ide jurnal SINTA 4 sekaligus aplikasi AI untuk toko pakaian "
        "(disesuaikan dengan implementasi backend saat ini)"
    )
    sub.runs[0].italic = True

    # --- Gambaran Fitur
    doc.add_heading("Gambaran Fitur", level=1)
    doc.add_paragraph(
        "Fitur utama aplikasi adalah merekomendasikan pakaian yang sebaiknya "
        "dibeli pengguna berdasarkan budget, foto pengguna atau mannequin, dan "
        "tujuan/gaya pemakaian. Yang membuatnya lebih kuat dari sekadar chatbot "
        "rekomendasi adalah hasilnya dibatasi budget dan dijamin hanya memakai "
        "produk yang benar-benar ada di katalog toko, sehingga masalahnya jelas "
        "dan bisa diukur untuk keperluan jurnal."
    )

    # --- Manfaat
    doc.add_heading("Manfaat bagi Pelanggan dan Pemilik Toko", level=1)
    doc.add_paragraph(
        "Aplikasi ini melayani dua pihak sekaligus: pelanggan terbantu memilih "
        "pakaian, pemilik toko terbantu menjual."
    )
    t = doc.add_table(rows=1, cols=3)
    t.style = "Light Grid Accent 1"
    hdr = t.rows[0].cells
    hdr[0].text, hdr[1].text, hdr[2].text = (
        "Pihak", "Masalah yang dialami", "Bantuan dari aplikasi"
    )
    manfaat = [
        ("Pelanggan", "Bingung memilih pakaian yang pas untuk acara tertentu",
         "Paket outfit sesuai acara/gaya, diambil dari katalog toko"),
        ("Pelanggan", "Takut belanja melebihi budget",
         "Total harga rekomendasi dijamin tidak melewati budget"),
        ("Pelanggan", "Tidak yakin pakaian cocok saat dipakai",
         "Pratinjau outfit pada manekin atau foto sendiri"),
        ("Pemilik toko", "Pelanggan pergi tanpa membeli karena bingung",
         "Rekomendasi mengarahkan pelanggan ke produk yang tersedia"),
        ("Pemilik toko", "Pelanggan hanya membeli satu item",
         "Paket outfit mendorong pembelian beberapa item sekaligus"),
        ("Pemilik toko", "Tidak tahu produk apa yang dicari pelanggan",
         "Data budget dan gaya pelanggan menjadi bahan analisis permintaan"),
    ]
    for row in manfaat:
        cells = t.add_row().cells
        for i, val in enumerate(row):
            cells[i].text = val

    # --- Cara Kerja
    doc.add_heading("Cara Kerja", level=1)
    doc.add_paragraph("Input dari pengguna:")
    add_bullets(doc, [
        "Budget total, misalnya Rp500.000",
        "Foto pengguna atau mannequin (jpg/png/webp)",
        "Prompt tambahan opsional berupa teks bebas, misalnya "
        "\"gaya formal untuk wawancara\" atau \"kasual, warna netral\"",
        "Mode visualisasi: manekin atau foto saya",
    ])
    doc.add_paragraph(
        "Output: satu paket outfit (atasan, bawahan, alas kaki, aksesori jika "
        "tersedia) dari katalog toko dengan total harga tidak melebihi budget, "
        "alasan singkat per item, penjelasan keseluruhan, dan sebuah gambar "
        "pratinjau outfit pada manekin atau foto pengguna."
    )

    # --- Arsitektur
    doc.add_heading("Arsitektur yang Diimplementasikan", level=1)
    doc.add_paragraph(
        "Sistem memisahkan tahap-tahap kecerdasan (berbasis Gemini) dari tahap "
        "visualisasi, sehingga mudah dikembangkan dan dievaluasi:"
    )
    add_bullets(doc, [
        "Analisis foto pengguna (Gemini Vision). Model multimodal membaca foto "
        "pengguna atau mannequin untuk menangkap konteks seperti tone warna, "
        "yang memperkaya proses pemilihan.",
        "Pemilihan produk (Gemini Flash). Model menerima seluruh katalog dari "
        "database beserta budget dan preferensi, lalu memilih kombinasi outfit "
        "yang paling cocok. Model hanya boleh memilih id produk yang ada; id "
        "karangan otomatis dibuang.",
        "Penjelasan (Gemini Flash). Model menghasilkan alasan berbahasa alami, "
        "misalnya \"kemeja ini cocok untuk wawancara karena ...\", dan "
        "menangani input bebas seperti \"rapi tapi tidak kaku\".",
        "Visualisasi outfit (lokal, Pillow). Produk terpilih dikomposisikan ke "
        "siluet manekin atau ke foto pengguna per kategori. Tahap ini berjalan "
        "offline tanpa biaya, sehingga tidak bergantung pada kuota model gambar.",
    ])
    doc.add_paragraph(
        "Jaminan konsistensi: harga dan total selalu dihitung dari data "
        "database, bukan dari keluaran model. Jika layanan AI sedang tidak "
        "tersedia, pemilihan jatuh ke fallback deterministik berbasis budget "
        "(satu item per kategori termurah yang muat budget) sehingga sistem "
        "tetap mengembalikan outfit beserta pratinjaunya."
    )
    doc.add_paragraph(
        "Kenapa bukan LLM saja? LLM murni sering melanggar budget atau mengarang "
        "produk yang tidak ada. Dengan katalog nyata sebagai satu-satunya sumber "
        "pilihan, perhitungan harga dari database, validasi id, dan fallback "
        "budget, hasilnya konsisten dan hanya memakai produk yang benar-benar "
        "ada. Perbedaan ini bisa dijadikan kontribusi utama jurnal."
    )

    # --- Stack teknis
    doc.add_heading("Stack Teknis", level=1)
    add_bullets(doc, [
        "Python 3.12, FastAPI, Uvicorn",
        "PostgreSQL dengan tabel category dan item (SQLAlchemy). Tabel category "
        "menyimpan name dan body_zone agar kategori konsisten dan pratinjau "
        "manekin akurat.",
        "Gemini API (google-genai) untuk analisis foto pengguna, pemilihan "
        "produk, dan penjelasan rekomendasi.",
        "Pillow untuk komposit visual outfit (manekin / foto)",
        "Scalar API Reference sebagai dokumentasi interaktif di endpoint /docs",
    ])
    doc.add_paragraph("Endpoint utama:")
    add_bullets(doc, [
        "POST /api/v1/categories — membuat kategori (name, body_zone)",
        "GET /api/v1/categories — melihat daftar kategori",
        "POST /api/v1/items — membuat produk baru dengan upload gambar",
        "GET /api/v1/items — melihat katalog",
        "POST /api/v1/recommendations — menghasilkan rekomendasi outfit + pratinjau",
    ])

    # --- Rumusan Masalah
    doc.add_heading("Contoh Rumusan Masalah Jurnal", level=1)
    q = doc.add_paragraph(
        "Bagaimana merancang sistem rekomendasi pakaian berbasis budget dan "
        "kebutuhan acara dengan memanfaatkan model AI multimodal (Gemini: "
        "analisis foto, pemilihan produk, dan penjelasan) yang dibatasi katalog "
        "nyata, sehingga membantu pelanggan memilih pakaian sesuai kebutuhan dan "
        "membantu pemilik toko meningkatkan penjualan, serta seberapa baik "
        "kepatuhan budget dan validitas produknya dibanding LLM murni?",
        style="Quote",
    )
    _ = q

    # --- Evaluasi
    doc.add_heading("Evaluasi yang Terukur", level=1)
    doc.add_paragraph(
        "Untuk SINTA 4, perbandingan pendekatan terbatas-katalog ini dengan LLM "
        "murni, ditambah uji responden, sudah cukup kuat."
    )
    t2 = doc.add_table(rows=1, cols=2)
    t2.style = "Light Grid Accent 1"
    h2 = t2.rows[0].cells
    h2[0].text, h2[1].text = "Yang diukur", "Cara ukurnya"
    evaluasi = [
        ("Kepatuhan budget",
         "Persentase rekomendasi yang tidak melebihi budget "
         "(sistem terbatas-katalog vs LLM murni)"),
        ("Validitas produk", "Persentase item yang benar-benar ada di katalog"),
        ("Relevansi dengan acara/gaya", "Penilaian 20 sampai 30 responden (skala 1 sampai 5)"),
        ("Kualitas pratinjau outfit", "Penilaian responden terhadap kejelasan visualisasi"),
        ("Kepuasan pengguna", "Kuesioner SUS"),
    ]
    for row in evaluasi:
        cells = t2.add_row().cells
        cells[0].text, cells[1].text = row

    # --- Evaluasi pemilik toko
    doc.add_heading("Evaluasi dari Sisi Pemilik Toko", level=1)
    add_bullets(doc, [
        "Wawancara atau kuesioner singkat dengan pemilik toko tentang kegunaan rekomendasi",
        "Persentase rekomendasi yang hanya memakai produk yang tersedia di toko",
        "Perbandingan sebelum dan sesudah memakai aplikasi, misalnya jumlah item per transaksi",
    ])

    # --- Yang perlu disiapkan
    doc.add_heading("Yang Perlu Disiapkan", level=1)
    add_bullets(doc, [
        "Daftar kategori produk (mis. atasan, bawahan, alas kaki, luaran, "
        "aksesori) beserta body_zone untuk posisi pada pratinjau",
        "Katalog produk minimal sekitar 50 sampai 100 item dengan atribut: "
        "nama, kategori (referensi), harga, gambar, dan deskripsi. Sebutkan "
        "warna produk di nama atau deskripsi agar AI dapat mempertimbangkannya.",
        "Gambar produk disimpan sebagai file di folder public/, database "
        "menyimpan path relatifnya",
        "Skenario uji, misalnya 20 kombinasi budget, foto, dan gaya berbeda",
        "Keputusan sumber data katalog: dari toko sungguhan atau dibuat sendiri dulu",
        "Gemini API key yang aktif disimpan di environment variable",
    ])

    # --- Catatan keterbatasan
    doc.add_heading("Catatan Keterbatasan", level=1)
    add_bullets(doc, [
        "Model gambar Gemini (Nano Banana) tidak memiliki kuota pada free tier, "
        "sehingga visualisasi dibuat secara lokal dengan Pillow (komposit "
        "manekin/foto), bukan virtual try-on fotorealistis.",
        "Virtual try-on fotorealistis tetap memungkinkan sebagai pengembangan "
        "lanjutan bila kuota/billing model gambar tersedia; arsitektur sudah "
        "memisahkan tahap visualisasi sehingga mudah ditukar.",
    ])

    return doc


if __name__ == "__main__":
    document = build()
    document.save(str(OUTPUT))
    print(f"Saved: {OUTPUT}")
